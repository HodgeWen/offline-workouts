#!/usr/bin/env python3
"""从 健身资料.xlsx 抽取预置数据，输出到 app/assets/seed/。

可重复执行；对合并单元格与脏数据做容错。
"""

from __future__ import annotations

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("需要 openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "健身资料.xlsx"
OUT = ROOT / "app" / "assets" / "seed"

SPLIT_SHEETS = {
    "gym_3day": "20训练计划-健身房三分化",
    "gym_4day_shoulder": "21训练计划-健身房四分化计划（肩单练版）",
    "gym_4day_arm": "22训练计划-健身房四分化计划（手臂单练版）",
    "home_3day": "23训练计划-居家健身",
}

DIET_SLOTS = [
    ("cut", "after_breakfast_early", "1减脂-早饭后练（早起版）", "早饭后练（早起）"),
    ("cut", "after_breakfast_late", "2减脂-早饭后练（晚起版）", "早饭后练（晚起）"),
    ("cut", "before_lunch", "3减脂-午饭前练", "午饭前练"),
    ("cut", "after_lunch", "4减脂-午饭后练", "午饭后练"),
    ("cut", "before_dinner", "5减脂-晚饭前练", "晚饭前练"),
    ("cut", "after_dinner", "6减脂-晚饭后练", "晚饭后练"),
    ("cut", "night", "7减脂-夜里练", "夜里练"),
    ("cut", "no_strength", "8减脂-无力训者", "无力训"),
    ("bulk", "after_breakfast_early", "9增肌-早饭后练（早起版）", "早饭后练（早起）"),
    ("bulk", "after_breakfast_late", "10增肌-早饭后练（晚起版）", "早饭后练（晚起）"),
    ("bulk", "before_lunch", "11增肌-午饭前练", "午饭前练"),
    ("bulk", "after_lunch", "12增肌-午饭后练", "午饭后练"),
    ("bulk", "before_dinner", "13增肌-晚饭前练", "晚饭前练"),
    ("bulk", "after_dinner", "14增肌-晚饭后练", "晚饭后练"),
    ("bulk", "night", "15增肌-夜里练", "夜里练"),
]


def cell(ws, r: int, c: int):
    v = ws.cell(r, c).value
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip()
    return v


def clean_text(s) -> str | None:
    if s is None:
        return None
    t = str(s).strip()
    t = re.sub(r"\s+", " ", t.replace("\n", " "))
    return t or None


def write_json(name: str, data) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {path.relative_to(ROOT)} ({_count(data)} items)")


def _count(data) -> str:
    if isinstance(data, list):
        return str(len(data))
    if isinstance(data, dict):
        if "items" in data and isinstance(data["items"], list):
            return str(len(data["items"]))
        return str(len(data))
    return "?"


def parse_sets_rule(text: str | None) -> dict:
    """解析『选1-2个动作 总共6-8组』一类文案。"""
    if not text:
        return {"pickMin": 1, "pickMax": 1, "setsMin": 3, "setsMax": 4, "raw": ""}
    raw = clean_text(text) or ""
    pick_min, pick_max = 1, 1
    sets_min, sets_max = 3, 4
    m = re.search(r"选\s*(\d+)\s*[-~～到至]\s*(\d+)\s*个", raw)
    if m:
        pick_min, pick_max = int(m.group(1)), int(m.group(2))
    else:
        m = re.search(r"选\s*(\d+)\s*个", raw)
        if m:
            pick_min = pick_max = int(m.group(1))
    m = re.search(r"总共\s*(\d+)\s*[-~～到至]\s*(\d+)\s*组", raw)
    if m:
        sets_min, sets_max = int(m.group(1)), int(m.group(2))
    else:
        m = re.search(r"总共\s*(\d+)\s*组", raw)
        if m:
            sets_min = sets_max = int(m.group(1))
    return {
        "pickMin": pick_min,
        "pickMax": pick_max,
        "setsMin": sets_min,
        "setsMax": sets_max,
        "raw": raw,
    }


def extract_split_templates(wb) -> tuple[list, list]:
    """返回 (templates, exercises_dedup)。"""
    templates = []
    exercises: OrderedDict[str, dict] = OrderedDict()

    for tid, sheet_name in SPLIT_SHEETS.items():
        ws = wb[sheet_name]
        days = []
        current_day = None
        current_slot = None
        current_region = None

        for r in range(1, ws.max_row + 1):
            c2 = clean_text(cell(ws, r, 2))
            c3 = clean_text(cell(ws, r, 3))
            c4 = clean_text(cell(ws, r, 4))
            c5 = clean_text(cell(ws, r, 5))
            c6 = clean_text(cell(ws, r, 6))

            if c2 and re.match(r"^Day\d+", c2, re.I):
                if current_day and current_slot:
                    current_day["slots"].append(current_slot)
                    current_slot = None
                if current_day:
                    days.append(current_day)
                title = c2
                parts = title.split("：", 1)
                regions = []
                if len(parts) > 1:
                    regions = [p.strip() for p in re.split(r"[+＋]", parts[1]) if p.strip()]
                current_day = {
                    "id": f"{tid}_d{len(days)+1}",
                    "title": title,
                    "regions": regions,
                    "slots": [],
                }
                current_region = None
                continue

            if not current_day:
                continue

            # 新部位行：B 列有部位名（非表头）
            if c2 and c2 not in ("部位",) and not c2.startswith("为什么") and not c2.startswith("练"):
                # 排除说明行
                if "会有" in c2 or "所以" in c2:
                    pass
                elif len(c2) <= 20 and ("选" in (c3 or "") or c4):
                    if current_slot:
                        current_day["slots"].append(current_slot)
                    current_region = c2.replace("\n", "").strip()
                    current_slot = {
                        "region": current_region,
                        "setsRule": parse_sets_rule(c3),
                        "exercises": [],
                    }
                    if c4 and c4 not in ("动作",):
                        _add_ex(exercises, current_slot, c4, current_region, c5, c6, tid)
                    continue

            # 延续行：仅有动作名
            if current_slot and c4 and c4 not in ("动作",):
                # 也可能是子部位（如 中胸 / 下胸）写在 C 列
                if c3 and ("选" in c3 or "组" in c3) and not c2:
                    # 子部位新 slot
                    current_day["slots"].append(current_slot)
                    sub = clean_text(c3.split("选")[0]) or current_region or "未命名"
                    current_slot = {
                        "region": f"{current_region or ''}/{sub}".strip("/"),
                        "setsRule": parse_sets_rule(c3),
                        "exercises": [],
                    }
                _add_ex(exercises, current_slot, c4, current_region or current_slot["region"], c5, c6, tid)

        if current_day:
            if current_slot:
                current_day["slots"].append(current_slot)
            days.append(current_day)

        place = "home" if tid.startswith("home") else "gym"
        split = "3day" if "3day" in tid else "4day"
        templates.append(
            {
                "id": tid,
                "name": sheet_name.split("-", 1)[-1],
                "place": place,
                "split": split,
                "source": "excel",
                "days": days,
            }
        )

    return templates, list(exercises.values())


def _add_ex(exercises, slot, name, region, joint_shoulder, joint_elbow, template_id):
    name = clean_text(name)
    if not name or name in ("动作",):
        return
    joints = []
    for j in (joint_shoulder, joint_elbow):
        jj = clean_text(j)
        if jj and jj not in ("肩关节", "肘关节", "关节活动"):
            joints.append(jj)
    entry = {
        "name": name,
        "region": region,
        "joints": joints,
        "equipmentTags": _guess_equipment(name),
        "templateIds": [template_id],
        "source": "excel",
    }
    slot["exercises"].append({"name": name, "joints": joints})
    key = name
    if key in exercises:
        existing = exercises[key]
        for t in entry["templateIds"]:
            if t not in existing["templateIds"]:
                existing["templateIds"].append(t)
        for j in joints:
            if j not in existing["joints"]:
                existing["joints"].append(j)
        if region and region not in (existing.get("regions") or []):
            existing.setdefault("regions", [existing.get("region")] if existing.get("region") else [])
            if region not in existing["regions"]:
                existing["regions"].append(region)
    else:
        entry["regions"] = [region] if region else []
        exercises[key] = entry


def _guess_equipment(name: str) -> list[str]:
    tags = []
    mapping = [
        ("杠铃", "杠铃"),
        ("哑铃", "哑铃"),
        ("史密斯", "史密斯"),
        ("器械", "器械"),
        ("蝴蝶机", "器械"),
        ("龙门架", "绳索"),
        ("绳索", "绳索"),
        ("弹力带", "弹力带"),
        ("引体", "自重"),
        ("俯卧撑", "自重"),
        ("双杠", "自重"),
    ]
    for k, tag in mapping:
        if k in name and tag not in tags:
            tags.append(tag)
    if not tags:
        tags.append("综合")
    return tags


def extract_foods(wb) -> list:
    ws = wb["19日常食物营养率"]
    foods = []
    category = None
    # 表结构不规整：碳水段 / 蛋白质段 / 脂肪段分列；先扫整表取「大类/食物/率」
    # 简化：遍历所有行，找食物名 + 数值率
    for r in range(1, ws.max_row + 1):
        c2 = clean_text(cell(ws, r, 2))
        c3 = clean_text(cell(ws, r, 3))
        c4 = cell(ws, r, 4)
        c5 = clean_text(cell(ws, r, 5))
        c6 = clean_text(cell(ws, r, 6))

        if c2 and c2 in ("碳水", "蛋白质", "脂肪", "其他"):
            # 大分区
            pass
        if c2 and c3 is None and c4 is None:
            # 可能是分区标题
            continue
        if c2 and c3 is None and not isinstance(c4, (int, float)):
            # 大类标题行如「米类主食」
            if c2 not in ("大类", "碳水", "蛋白质", "脂肪"):
                category = c2
            continue
        if c2 and c3 and c2 not in ("大类",):
            category = c2
        name = c3
        if not name or name in ("食物",):
            continue
        rate = None
        if isinstance(c4, (int, float)):
            rate = float(c4)
        elif isinstance(c4, str):
            try:
                rate = float(c4)
            except ValueError:
                rate = None
        if rate is None:
            continue
        # 粗分宏量类型
        macro = "carb"
        if category and ("蛋白" in category or "肉" in category or "蛋" in category or "奶" in category):
            macro = "protein"
        if category and ("脂肪" in category or "坚果" in category or "油" in category):
            macro = "fat"
        # 用附近表头推断：sheet 后半段是蛋白质
        foods.append(
            {
                "name": name,
                "category": category or "未分类",
                "rate": rate,
                "macroHint": macro,
                "gi": c5,
                "note": c6,
                "source": "excel",
            }
        )

    # 二次扫描：标记蛋白质区（常见从某行开始）
    # 改进：若 category 含「蛋白」则已标；否则按名称启发式
    protein_kw = ("肉", "蛋", "奶", "蛋白粉", "豆腐", "虾", "鱼", "鸡", "牛", "猪", "羊")
    fat_kw = ("坚果", "油", "牛油果", "芝麻")
    for f in foods:
        n = f["name"]
        cat = f["category"]
        if any(k in cat for k in ("蛋白", "肉蛋", "乳")) or any(k in n for k in protein_kw):
            if f["macroHint"] == "carb" and f["rate"] < 0.5:
                f["macroHint"] = "protein"
        if any(k in cat for k in ("脂肪",)) or any(k in n for k in fat_kw):
            if "面包" not in n:
                f["macroHint"] = "fat"
    return foods


def extract_cardio(wb) -> list:
    ws = wb["16有氧热量消耗"]
    items = []
    in_method2 = False
    category = None
    for r in range(1, ws.max_row + 1):
        c2 = clean_text(cell(ws, r, 2))
        c3 = clean_text(cell(ws, r, 3))
        c4 = cell(ws, r, 4)
        c5 = cell(ws, r, 5)
        if c2 and "方法二" in c2:
            in_method2 = True
        if c2 and c2.startswith("【数据来源】"):
            break
        if not in_method2:
            continue
        if c3 and not c3.startswith("*"):
            category = c3.split("*")[0].strip()
        detail = clean_text(c4) if not isinstance(c4, (int, float)) else str(c4)
        if isinstance(c4, (int, float)):
            # 速度/功率类数值
            unit = "km/h" if category and ("跑" in category or "骑" in category) else ""
            detail = f"{c4}{unit}".rstrip()
        if not isinstance(c5, (int, float)):
            continue
        per_kg = float(c5)
        if per_kg <= 0 or per_kg > 30:
            continue
        name_parts = [p for p in (category, detail) if p]
        name = " / ".join(name_parts) if name_parts else None
        if not name:
            continue
        items.append(
            {
                "name": name,
                "category": category,
                "kcalPerKgPerHour": per_kg,
                "source": "excel",
            }
        )
    seen = set()
    out = []
    for it in items:
        if it["name"] in seen:
            continue
        seen.add(it["name"])
        out.append(it)
    return out


def extract_faq(wb) -> list:
    items = []
    for goal, sheet in (("cut", "17减脂-问答汇总"), ("bulk", "18增肌-问答汇总")):
        ws = wb[sheet]
        # 先收集目录问题，再匹配正文
        catalog = []
        for r in range(1, min(ws.max_row, 40) + 1):
            text = clean_text(cell(ws, r, 2))
            if text and re.match(r"^\d+\.", text) and "？" in text and len(text) < 100:
                catalog.append(text)

        pending_q = None
        for r in range(1, ws.max_row + 1):
            raw = cell(ws, r, 2)
            if raw is None:
                continue
            raw_s = str(raw)
            text = clean_text(raw)
            if not text:
                continue
            # 正文问题行：含数字.与？，且单元格多行或前后空白
            m = re.search(r"(\d+\.[^？\n]*？)", raw_s, re.S)
            if m and len(m.group(1)) < 120:
                # 跳过纯目录（答案在后面更长段落）
                q = clean_text(m.group(1))
                # 若整格几乎只有问题，视为问题行
                if len(text) < 120:
                    pending_q = q
                    continue
            if pending_q and len(text) > 50 and not re.match(r"^\d+\.", text):
                items.append(
                    {
                        "goal": goal,
                        "question": pending_q,
                        "answer": text if len(text) < 4000 else text[:4000] + "…",
                        "source": "excel",
                    }
                )
                pending_q = None
    return items


def extract_anatomy(wb) -> dict:
    ws = wb["27健身解剖总结（文字版）"]
    table_a = []
    table_b = []
    mode = None
    current_joint = None

    # A 表
    for r in range(1, 40):
        c2 = clean_text(cell(ws, r, 2))
        c3 = clean_text(cell(ws, r, 3))
        if c2 and "A表" in c2:
            mode = "A"
            continue
        if c2 and "B表" in c2:
            break
        if mode != "A":
            continue
        if c2 and c2 not in ("关节活动",) and c3 in (
            "屈",
            "伸",
            "外展",
            "内收",
            "水平外展",
            "水平内收",
            "足跖屈",
            "上提",
            "下沉",
            "上旋",
            "下旋",
        ):
            current_joint = c2.split("（")[0].strip()
            action = c3
        elif c3 in (
            "屈",
            "伸",
            "外展",
            "内收",
            "水平外展",
            "水平内收",
            "足跖屈",
            "上提",
            "下沉",
            "上旋",
            "下旋",
        ):
            action = c3
        else:
            continue
        if not current_joint:
            continue
        desc = clean_text(cell(ws, r, 4))
        example = clean_text(cell(ws, r, 5))
        muscles = []
        for c in range(6, 11):
            m = clean_text(cell(ws, r, c))
            if m:
                # 去掉脚注数字
                m = re.sub(r"\d+$", "", m)
                muscles.append(m)
        table_a.append(
            {
                "joint": current_joint,
                "action": action,
                "description": desc,
                "example": example,
                "muscles": muscles,
                "source": "excel",
            }
        )

    # B 表：肌肉 → 关节活动/示例动作（简化抽取带 √ 的单元格）
    headers = []
    section_joint = None
    for r in range(40, ws.max_row + 1):
        c2 = clean_text(cell(ws, r, 2))
        row_vals = [clean_text(cell(ws, r, c)) for c in range(2, 12)]
        if c2 and ("练" in c2 and "：") or (c2 and c2.startswith("练")):
            continue
        if c2 == "关节活动":
            headers = [clean_text(cell(ws, r, c)) for c in range(5, 12)]
            continue
        if c2 and c2.endswith("关节"):
            section_joint = c2
        action = clean_text(cell(ws, r, 3))
        if action not in (
            "屈",
            "伸",
            "外展",
            "内收",
            "水平外展",
            "水平内收",
            "足跖屈",
        ):
            continue
        if not section_joint:
            continue
        for i, h in enumerate(headers):
            if not h:
                continue
            val = clean_text(cell(ws, r, 5 + i))
            if not val or "√" not in val:
                continue
            example = val.replace("√", "").strip()
            example = re.sub(r"^\d+\s*", "", example)
            table_b.append(
                {
                    "muscle": h,
                    "joint": section_joint,
                    "action": action,
                    "example": example or None,
                    "source": "excel",
                }
            )

    return {"tableA": table_a, "tableB": table_b, "source": "excel"}


def build_diet_rules() -> dict:
    """固化热量/宏量规则（对齐原表系数；配额取可执行默认值）。"""
    meal_templates = {}
    for goal, slot, _sheet, label in DIET_SLOTS:
        meal_templates.setdefault(goal, {})[slot] = {
            "id": f"{goal}_{slot}",
            "label": label,
            "trainingDayMeals": _default_meals(slot, training=True),
            "restDayMeals": _default_meals(slot, training=False),
        }

    return {
        "version": 1,
        "mifflin": {
            "male": "weight*9.99 + heightCm*6.25 - age*4.92 + 5",
            "female": "weight*9.99 + heightCm*6.25 - age*4.92 - 161",
        },
        "tdeeFactor": 0.7,  # 无运动总消耗 = BMR / 0.7
        "intakeFactor": {"cut": 0.64, "bulk": 0.84},
        "strengthKcal": {
            "male": {"beginner": 150, "intermediate": 200, "advanced": 250},
            "female": {"beginner": 100, "intermediate": 150, "advanced": 200},
        },
        # 配额 g/kg（原表需用户输入；此处固化可执行默认，对齐常见用法）
        "macroPerKg": {
            "cut": {
                "training": {"carb": 3.0, "protein": 2.0},
                "rest": {"carb": 2.0, "protein": 2.0},
            },
            "bulk": {
                "training": {"carb": 5.0, "protein": 2.0},
                "rest": {"carb": 4.0, "protein": 2.0},
            },
        },
        "fatGuideGrams": {
            "cut": {"male": 60, "female": 50},
            "bulk": {"male": 80, "female": 70},
        },
        "cardioAdvice": {
            "bulkDefault": False,
            "cut": {
                "above80kg": False,
                "between70and80": "optional_if_hungry",
                "below70kg": {"hoursPerWeek": 2},
            },
        },
        "mealTemplates": meal_templates,
        "source": "excel|curated",
    }


def _default_meals(slot: str, training: bool) -> list:
    """按训练时段给出餐序摘要。"""
    if not training:
        return [
            {"name": "早饭", "role": "other", "carbShare": 0.25, "proteinShare": 0.25, "note": "鸡蛋+牛奶打底"},
            {"name": "午饭", "role": "other", "carbShare": 0.35, "proteinShare": 0.35, "note": "主食+瘦肉菜"},
            {"name": "晚饭", "role": "other", "carbShare": 0.30, "proteinShare": 0.30, "note": "主食+瘦肉菜"},
            {"name": "零食/夜宵", "role": "snack", "carbShare": 0.10, "proteinShare": 0.10, "note": "少量，可挪到正餐"},
        ]

    # 训练日：按时段排列 练前/练后/其他
    catalogs = {
        "after_breakfast_early": ["早饭=练前餐", "练后餐", "午饭=其他餐", "晚饭=其他餐", "零食/夜宵"],
        "after_breakfast_late": ["早饭=练前餐", "练后餐", "午饭=其他餐", "晚饭=其他餐", "零食/夜宵"],
        "before_lunch": ["早饭=其他餐", "练前加餐", "午饭=练后餐", "晚饭=其他餐", "零食/夜宵"],
        "after_lunch": ["早饭=其他餐", "午饭=练前餐", "练后餐", "晚饭=其他餐", "零食/夜宵"],
        "before_dinner": ["早饭=其他餐", "午饭=其他餐", "练前加餐", "晚饭=练后餐", "零食/夜宵"],
        "after_dinner": ["早饭=其他餐", "午饭=其他餐", "晚饭=练前餐", "练后餐", "零食/夜宵"],
        "night": ["早饭=其他餐", "午饭=其他餐", "晚饭=其他餐", "练前加餐", "练后餐/夜宵"],
        "no_strength": ["早饭", "午饭", "晚饭", "零食/夜宵"],
    }
    names = catalogs.get(slot, catalogs["after_breakfast_early"])
    meals = []
    for n in names:
        role = "other"
        if "练前" in n:
            role = "pre"
        elif "练后" in n:
            role = "post"
        elif "零食" in n or "夜宵" in n:
            role = "snack"
        share_c = 0.35 if role == "post" else 0.2 if role == "pre" else 0.15 if role == "snack" else 0.25
        share_p = 0.30 if role == "post" else 0.2 if role == "pre" else 0.15 if role == "snack" else 0.25
        meals.append(
            {
                "name": n,
                "role": role,
                "carbShare": share_c,
                "proteinShare": share_p,
                "note": "主食+瘦肉按克数等价互换（见今日饮食方案）",
            }
        )
    # 归一化 share
    cs = sum(m["carbShare"] for m in meals) or 1
    ps = sum(m["proteinShare"] for m in meals) or 1
    for m in meals:
        m["carbShare"] = round(m["carbShare"] / cs, 3)
        m["proteinShare"] = round(m["proteinShare"] / ps, 3)
    return meals


def enrich_exercises_with_anatomy(exercises: list, anatomy: dict) -> list:
    """用 A 表关节活动补全关联肌肉文案。"""
    # action -> rows（同一动作名可能跨关节，取全部）
    by_action: dict[str, list] = {}
    for row in anatomy.get("tableA", []):
        by_action.setdefault(row["action"], []).append(row)

    action_keys = sorted(by_action.keys(), key=len, reverse=True)

    for ex in exercises:
        muscles = []
        notes = []
        for j in ex.get("joints") or []:
            # j 可能是「[宽握]肩内收」或「肩内收mix肩伸」
            text = re.sub(r"\[.*?\]", "", j).strip()
            parts = re.split(r"mix|/|、", text)
            for part in parts:
                part = part.strip()
                if not part:
                    continue
                matched = None
                # 优先匹配「肩内收」整段；否则找最长 action 后缀
                for ak in action_keys:
                    if part.endswith(ak) or ak in part:
                        # 若带关节前缀则过滤
                        cands = by_action[ak]
                        if "肩" in part:
                            cands = [r for r in cands if "肩" in r["joint"]] or cands
                        elif "肘" in part:
                            cands = [r for r in cands if "肘" in r["joint"]] or cands
                        elif "髋" in part:
                            cands = [r for r in cands if "髋" in r["joint"]] or cands
                        elif "膝" in part:
                            cands = [r for r in cands if "膝" in r["joint"]] or cands
                        matched = cands[0]
                        break
                if matched:
                    for m in matched["muscles"]:
                        if m not in muscles:
                            muscles.append(m)
                    if matched.get("description"):
                        note = f"{matched['joint']}{matched['action']}：{matched['description']}"
                        if note not in notes:
                            notes.append(note)
        ex["muscles"] = muscles
        if notes:
            ex["anatomyNote"] = "；".join(notes[:2])
        elif muscles:
            ex["anatomyNote"] = f"主要目标肌：{'、'.join(muscles[:4])}"
            ex["source"] = "excel|curated"
        else:
            # 按部位给常识补全
            region = (ex.get("region") or "").split("/")[0]
            curated = {
                "背": ("背阔肌、斜方肌、菱形肌", "肩伸/肩内收/肩水平外展相关拉类动作"),
                "胸": ("胸大肌、肩前束、肱三头肌", "肩水平内收与肘伸推类动作"),
                "肩后束": ("三角肌后束、冈下肌", "肩水平外展为主"),
                "肩前中束": ("三角肌前束/中束", "肩屈与肩外展"),
                "肱二头": ("肱二头肌、肱肌", "肘屈"),
                "肱三头": ("肱三头肌", "肘伸"),
                "臀腿": ("股四头肌、腘绳肌、臀大肌", "膝伸/髋伸为主"),
                "腿": ("股四头肌、腘绳肌、臀大肌", "膝伸/髋伸为主"),
            }
            for k, (ms, note) in curated.items():
                if k in region:
                    ex["muscles"] = ms.split("、")
                    ex["anatomyNote"] = note
                    ex["source"] = "excel|curated"
                    break
    return exercises


def main() -> int:
    if not XLSX.exists():
        print(f"找不到 {XLSX}", file=sys.stderr)
        return 1
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    templates, exercises = extract_split_templates(wb)
    anatomy = extract_anatomy(wb)
    exercises = enrich_exercises_with_anatomy(exercises, anatomy)
    foods = extract_foods(wb)
    cardio = extract_cardio(wb)
    faq = extract_faq(wb)
    diet_rules = build_diet_rules()

    write_json("split_templates.json", {"version": 1, "templates": templates})
    write_json("exercises.json", {"version": 1, "exercises": exercises})
    write_json("anatomy.json", {"version": 1, **anatomy})
    write_json("foods.json", {"version": 1, "foods": foods})
    write_json("cardio.json", {"version": 1, "items": cardio})
    write_json("faq.json", {"version": 1, "items": faq})
    write_json("diet_rules.json", diet_rules)
    write_json(
        "meta.json",
        {
            "seedVersion": 1,
            "sourceFile": "健身资料.xlsx",
            "generatedBy": "tools/import_excel.py",
        },
    )
    print("done.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
